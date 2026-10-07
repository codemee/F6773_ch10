"""Create individual Word certificates while preserving a supplied DOCX template."""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook


REQUIRED_COLUMNS = ("學員編號", "姓名", "Email", "課程名稱", "完成日期")
PLACEHOLDERS = {
    "{{姓名}}": "姓名",
    "{{課程名稱}}": "課程名稱",
    "{{完成日期}}": "完成日期",
}


@dataclass(frozen=True)
class Attendee:
    row_number: int
    student_id: str
    name: str
    email: str
    course: str
    completion_date: str


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def format_date(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y年%m月%d日")
    return clean(value)


def read_attendees(path: Path) -> tuple[list[Attendee], list[str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return [], ["Excel 沒有標題列。"]

    headers = {clean(value): index for index, value in enumerate(header_row) if clean(value)}
    missing_columns = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing_columns:
        return [], [f"Excel 缺少必要欄位：{', '.join(missing_columns)}。"]

    attendees: list[Attendee] = []
    issues: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        values = {name: row[index] if index < len(row) else None for name, index in headers.items()}
        # Completely empty trailing rows are ignored.
        if not any(clean(values.get(name)) for name in REQUIRED_COLUMNS):
            continue
        required_values = {
            "學員編號": clean(values["學員編號"]),
            "姓名": clean(values["姓名"]),
            "Email": clean(values["Email"]),
            "課程名稱": clean(values["課程名稱"]),
            "完成日期": format_date(values["完成日期"]),
        }
        blank_fields = [name for name, value in required_values.items() if not value]
        if blank_fields:
            issues.append(f"第 {row_number} 列：缺少 {', '.join(blank_fields)}。")
            continue
        attendees.append(
            Attendee(
                row_number=row_number,
                student_id=required_values["學員編號"],
                name=required_values["姓名"],
                email=required_values["Email"],
                course=required_values["課程名稱"],
                completion_date=required_values["完成日期"],
            )
        )

    duplicate_ids = {student_id for student_id, count in Counter(a.student_id for a in attendees).items() if count > 1}
    if duplicate_ids:
        for student_id in sorted(duplicate_ids):
            row_numbers = ", ".join(str(a.row_number) for a in attendees if a.student_id == student_id)
            issues.append(f"學員編號「{student_id}」重複（第 {row_numbers} 列）。")
        attendees = [a for a in attendees if a.student_id not in duplicate_ids]
    return attendees, issues


def replace_tokens_raw(content: bytes, replacements: dict[str, str]) -> tuple[bytes, int]:
    """Replace literal template tokens without parsing or rebuilding Word XML.

    Keeping all source bytes other than the token itself preserves Word's run
    properties, drawing markup, and compatibility-sensitive XML exactly.
    """
    changed = 0
    for token, value in replacements.items():
        token_bytes = token.encode("utf-8")
        count = content.count(token_bytes)
        if count:
            content = content.replace(token_bytes, value.encode("utf-8"))
            changed += count
    return content, changed


def fill_template(template: Path, destination: Path, attendee: Attendee) -> None:
    replacements = {
        "{{姓名}}": attendee.name,
        "{{課程名稱}}": attendee.course,
        "{{完成日期}}": attendee.completion_date,
    }
    with zipfile.ZipFile(template, "r") as source:
        word_xml = [source.read(item.filename) for item in source.infolist() if item.filename.startswith("word/") and item.filename.endswith(".xml")]
    missing_tokens = [token for token in replacements if not any(token.encode("utf-8") in content for content in word_xml)]
    if missing_tokens:
        raise ValueError("範本缺少欄位：" + "、".join(missing_tokens))

    changed_parts = 0
    with zipfile.ZipFile(template, "r") as source, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                content, replacements_in_part = replace_tokens_raw(content, replacements)
                changed_parts += replacements_in_part
            target.writestr(item, content)
    if not changed_parts:
        destination.unlink(missing_ok=True)
        raise ValueError("範本內找不到可替換的證書欄位。")


def safe_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(". ") or "未命名"


def write_report(path: Path, created: list[Path], issues: list[str]) -> None:
    lines = ["研習證書產生結果", "=" * 20, f"成功產生：{len(created)} 份"]
    lines.extend(f"- {file.name}" for file in created)
    lines.extend(["", f"需要處理：{len(issues)} 項"])
    lines.extend(f"- {issue}" for issue in issues) if issues else lines.append("- 無")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="從 Excel 名單及 Word 範本產生研習證書。")
    parser.add_argument("--roster", type=Path, default=Path("參加者名單.xlsx"), help="Excel 參加者名單")
    parser.add_argument("--template", type=Path, default=Path("研習證書範本.docx"), help="Word 證書範本")
    parser.add_argument("--output", type=Path, default=Path("輸出證書"), help="輸出資料夾")
    args = parser.parse_args()

    if not args.roster.is_file() or not args.template.is_file():
        missing = [str(path) for path in (args.roster, args.template) if not path.is_file()]
        print("找不到檔案：" + "、".join(missing), file=sys.stderr)
        return 2

    attendees, issues = read_attendees(args.roster)
    args.output.mkdir(parents=True, exist_ok=True)
    created: list[Path] = []
    for attendee in attendees:
        filename = f"{safe_filename(attendee.student_id)}_{safe_filename(attendee.name)}_研習證書.docx"
        destination = args.output / filename
        try:
            fill_template(args.template, destination, attendee)
            created.append(destination)
        except (OSError, ValueError) as error:
            issues.append(f"第 {attendee.row_number} 列：{error}")
    report = args.output / "產生報告.txt"
    write_report(report, created, issues)
    print(f"成功產生 {len(created)} 份證書。報告：{report}")
    if issues:
        print("需要處理的項目：")
        print("\n".join(f"- {issue}" for issue in issues))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
